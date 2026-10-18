"""Turn returned clinician GRBAS ratings into ``real_training_features.csv``.

This is the missing link between the clinician deliverable and the trainer:

    voiced_rating.html  ->  clinician ratings CSV  ->  [this script]  ->
    real_training_features.csv  ->  train_grbas.py --mode finetune

The rating page (``build_voiced_clinician_web.py``) and the Google-Sheet
workbook (``build_voiced_clinician_sheet.py``) hand back perceptual labels but
no acoustics; ``train_grbas.py`` wants one wide row per recording with the 167
columns in ``grbas_model.FEATURE_NAMES`` plus ``sample_id``, ``speaker_id`` and
``clinician_grbas_{g,r,b,a,s}``. This script joins the two: for every rated,
usable recording it locates the WAV, runs the same feature extractor the live
app runs, and writes the trainer's input table.

In VOICED one recording == one speaker, so ``speaker_id`` is the session id and
the trainer's speaker-grouped split is leakage-free by construction.

Usage
-----
    python build_real_training_features.py --ratings voiced_grbas_clinician_ratings.csv

    # QC only -- prints the label distribution without extracting features
    python build_real_training_features.py --ratings ratings.csv --dry-run

Both delivery routes are accepted: the rating page's CSV
(``session_id,diagnosis,group,G..S,usable,quality,comment``) and the workbook
(``sample_usable``/``sample_quality``/``clinician_comments``), as .csv or .xlsx.
"""

from __future__ import annotations

import argparse
import csv
from collections import Counter
from pathlib import Path
from typing import Dict, List, Optional, Tuple

BASE_DIR = Path(__file__).resolve().parent

DEFAULT_RATINGS = BASE_DIR / "voiced_grbas_clinician_ratings.csv"
DEFAULT_AUDIO_DIR = BASE_DIR.parent / "root" / "voice_app" / "static" / "voiced_audio"
DEFAULT_OUT = BASE_DIR / "real_training_features.csv"

GRBAS_DIMS = "GRBAS"

# VOICED ships 8 kHz recordings; the app analyses whatever the browser records,
# which is 44.1/48 kHz in practice. Features are therefore computed at one fixed
# rate so the training rows and the live rows describe the same bands. 16 kHz is
# also the floor for librosa's spectral_contrast defaults (fmin 200, 6 bands) --
# below it feature extraction raises "Frequency band exceeds Nyquist".
DEFAULT_TARGET_SR = 16000

# train_grbas.py refuses to write the production model unless the real test
# split has --min-real-test rows (default 15). With its 70/15/15 speaker-grouped
# split that means roughly this many rated recordings.
MIN_ROWS_FOR_PRODUCTION_MODEL = 100

# Column aliases: rating page (first) vs. clinician workbook (second).
ALIASES = {
    "session_id": ("session_id", "sample_id", "id"),
    "usable": ("usable", "sample_usable"),
    "quality": ("quality", "sample_quality"),
    "comment": ("comment", "clinician_comments", "comments"),
    "diagnosis": ("diagnosis",),
    "group": ("group",),
    "age": ("age",),
    "gender": ("gender",),
}

NO_VALUES = {"no", "n", "false", "0", "unusable", "לא"}


# ---------------------------------------------------------------------------
# Reading the clinician's file
# ---------------------------------------------------------------------------

def field(row: Dict[str, str], name: str) -> str:
    """Read a logical column, whichever of its aliases the file happens to use."""
    for key in ALIASES.get(name, (name,)):
        if key in row and row[key] is not None:
            return str(row[key]).strip()
    return ""


def load_ratings(path: Path) -> List[Dict[str, str]]:
    """Load the returned ratings from .csv or .xlsx into plain row dicts."""
    if not path.exists():
        raise SystemExit(f"ratings file not found: {path}")

    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        workbook = load_workbook(path, data_only=True)
        sheet = (workbook["GRBAS Ratings"] if "GRBAS Ratings" in workbook.sheetnames
                 else workbook.worksheets[0])
        rows = sheet.iter_rows(values_only=True)
        try:
            header = [str(c).strip() if c is not None else "" for c in next(rows)]
        except StopIteration:
            return []
        return [
            {header[i]: ("" if value is None else value)
             for i, value in enumerate(values) if i < len(header)}
            for values in rows
            if any(v not in (None, "") for v in values)
        ]

    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def read_labels(row: Dict[str, str]) -> Optional[Dict[str, int]]:
    """Return {G,R,B,A,S: 0-3}, or None if any dimension is blank/invalid."""
    labels: Dict[str, int] = {}
    for dim in GRBAS_DIMS:
        raw = ""
        for key in (dim, dim.lower(), f"clinician_grbas_{dim.lower()}"):
            if key in row and str(row[key]).strip() != "":
                raw = str(row[key]).strip()
                break
        if raw == "":
            return None
        try:
            value = int(round(float(raw)))
        except (TypeError, ValueError):
            return None
        if not 0 <= value <= 3:
            return None
        labels[dim] = value
    return labels


def triage(rows: List[Dict[str, str]]) -> Tuple[List[dict], Dict[str, List[str]]]:
    """Split the returned rows into usable rated recordings and dropped ones.

    Dropped rows are kept per reason so the QC report can name them instead of
    silently shrinking the training set.
    """
    kept: List[dict] = []
    dropped: Dict[str, List[str]] = {
        "no_session_id": [], "marked_unusable": [], "not_rated": [], "duplicate": [],
    }
    seen = set()

    for row in rows:
        session_id = field(row, "session_id")
        if not session_id:
            dropped["no_session_id"].append("<blank>")
            continue
        if field(row, "usable").lower() in NO_VALUES:
            dropped["marked_unusable"].append(session_id)
            continue
        labels = read_labels(row)
        if labels is None:
            dropped["not_rated"].append(session_id)
            continue
        if session_id in seen:
            dropped["duplicate"].append(session_id)
            continue
        seen.add(session_id)
        kept.append({
            "session_id": session_id,
            "labels": labels,
            "diagnosis": field(row, "diagnosis"),
            "group": field(row, "group"),
            "age": field(row, "age"),
            "gender": field(row, "gender"),
            "quality": field(row, "quality"),
            "comment": field(row, "comment"),
        })

    return kept, dropped


# ---------------------------------------------------------------------------
# Audio
# ---------------------------------------------------------------------------

def resolve_audio(session_id: str, audio_dir: Path, db_dir: Optional[Path]) -> Optional[Path]:
    """Locate the WAV for a session, rendering it from VOICED when missing."""
    wav = audio_dir / f"{session_id}.wav"
    if wav.exists():
        return wav
    if db_dir is not None:
        import build_voiced_clinician_sheet as sheet

        audio_dir.mkdir(parents=True, exist_ok=True)
        if sheet.convert_to_wav(session_id, db_dir, audio_dir) and wav.exists():
            return wav
    return None


def prepare_audio(wav: Path, target_sr: int, work_dir: Path) -> Path:
    """Return ``wav`` at ``target_sr``, resampling into ``work_dir`` if needed.

    Upsampling adds no real bandwidth to an 8 kHz recording, but it puts every
    row on the same frequency grid as the app's own recordings — without it the
    spectral features would be systematically shifted against production audio.
    """
    if not target_sr:
        return wav

    import soundfile as sf

    info = sf.info(str(wav))
    if info.samplerate == target_sr:
        return wav

    import librosa

    y, sr = librosa.load(str(wav), sr=target_sr, mono=True)
    work_dir.mkdir(parents=True, exist_ok=True)
    out = work_dir / f"{wav.stem}_{target_sr}.wav"
    sf.write(str(out), y, target_sr, subtype="PCM_16")
    return out


# ---------------------------------------------------------------------------
# QC report
# ---------------------------------------------------------------------------

def qc_report(kept: List[dict], dropped: Dict[str, List[str]], total_rows: int) -> None:
    print(f"rows in file            {total_rows}")
    for reason, ids in dropped.items():
        if ids:
            preview = ", ".join(ids[:6]) + (" ..." if len(ids) > 6 else "")
            print(f"  dropped: {reason:<16} {len(ids):>4}   {preview}")
    print(f"usable rated recordings {len(kept)}")

    if not kept:
        print("\nnothing to train on yet.")
        return

    print("\nlabel distribution (count per grade)")
    print(f"      {'0':>6}{'1':>6}{'2':>6}{'3':>6}   mean")
    for dim in GRBAS_DIMS:
        counts = Counter(row["labels"][dim] for row in kept)
        mean = sum(row["labels"][dim] for row in kept) / len(kept)
        cells = "".join(f"{counts.get(grade, 0):>6}" for grade in range(4))
        print(f"  {dim}   {cells}   {mean:.2f}")

    for dim in GRBAS_DIMS:
        counts = Counter(row["labels"][dim] for row in kept)
        if len(counts) < 2:
            print(f"  WARNING: {dim} has a single value across every recording — "
                  f"nothing to learn on that dimension.")

    if len(kept) < MIN_ROWS_FOR_PRODUCTION_MODEL:
        print(f"\nWARNING: {len(kept)} rated recordings. train_grbas.py --mode "
              f"finetune needs roughly {MIN_ROWS_FOR_PRODUCTION_MODEL} before its "
              f"real test split reaches --min-real-test (15), and will refuse to "
              f"write grbas_model.joblib below that.")


# ---------------------------------------------------------------------------
# Feature table
# ---------------------------------------------------------------------------

META_COLUMNS = [
    "sample_id", "speaker_id", "source_type",
    "diagnosis", "group", "age", "gender", "sample_quality",
] + [f"clinician_grbas_{dim.lower()}" for dim in GRBAS_DIMS]


def build_feature_table(kept: List[dict], audio_dir: Path, db_dir: Optional[Path],
                        out_path: Path,
                        target_sr: int = DEFAULT_TARGET_SR) -> Tuple[int, List[str]]:
    """Extract features for every kept recording and write the trainer's CSV."""
    import tempfile

    from grbas_model import FEATURE_NAMES, extract_features

    fieldnames = META_COLUMNS + list(FEATURE_NAMES)
    failures: List[str] = []
    written = 0
    work_dir = Path(tempfile.mkdtemp(prefix="grbas-resample-"))

    with out_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()

        for index, row in enumerate(kept, start=1):
            session_id = row["session_id"]
            wav = resolve_audio(session_id, audio_dir, db_dir)
            if wav is None:
                failures.append(f"{session_id}: no audio")
                continue
            try:
                features = extract_features(str(prepare_audio(wav, target_sr, work_dir)))
            except Exception as exc:                      # noqa: BLE001 - report, keep going
                failures.append(f"{session_id}: {type(exc).__name__}: {exc}")
                continue

            record = {
                "sample_id": session_id,
                # One recording == one speaker in VOICED, so the trainer's
                # speaker-grouped split cannot leak a speaker across splits.
                "speaker_id": session_id,
                "source_type": "real",
                "diagnosis": row["diagnosis"],
                "group": row["group"],
                "age": row["age"],
                "gender": row["gender"],
                "sample_quality": row["quality"],
            }
            for dim in GRBAS_DIMS:
                record[f"clinician_grbas_{dim.lower()}"] = row["labels"][dim]
            for name in FEATURE_NAMES:
                record[name] = features.get(name, 0.0)

            writer.writerow(record)
            written += 1

            if index % 20 == 0 or index == len(kept):
                print(f"  ... {index}/{len(kept)} processed ({written} written)",
                      flush=True)

    import shutil

    shutil.rmtree(work_dir, ignore_errors=True)
    return written, failures


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=__doc__.split("\n\n")[0],
        formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--ratings", type=Path, default=DEFAULT_RATINGS,
                        help="CSV/XLSX returned by the clinician")
    parser.add_argument("--audio-dir", type=Path, default=DEFAULT_AUDIO_DIR,
                        help="folder holding <session_id>.wav")
    parser.add_argument("--db-dir", type=Path, default=None,
                        help="VOICED database folder; used to render any WAV "
                             "missing from --audio-dir")
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT,
                        help="output feature table for train_grbas.py")
    parser.add_argument("--dry-run", action="store_true",
                        help="print the QC report only; extract nothing")
    parser.add_argument("--target-sr", type=int, default=DEFAULT_TARGET_SR,
                        help="resample every recording to this rate before "
                             "feature extraction (0 = keep native rate)")
    parser.add_argument("--limit", type=int, default=None,
                        help="process only the first N rated recordings "
                             "(smoke test)")
    return parser.parse_args()


def main() -> None:
    args = parse_args()

    rows = load_ratings(args.ratings)
    kept, dropped = triage(rows)
    qc_report(kept, dropped, len(rows))

    if args.limit is not None:
        kept = kept[:args.limit]
        print(f"\n--limit {args.limit}: processing {len(kept)} recording(s)")

    if args.dry_run:
        print("\n--dry-run: no features extracted, no file written.")
        return
    if not kept:
        raise SystemExit("no usable rated recordings — nothing to write.")

    rate = f"at {args.target_sr} Hz" if args.target_sr else "at their native rate"
    print(f"\nextracting features from {len(kept)} recording(s) {rate}...")
    written, failures = build_feature_table(kept, args.audio_dir, args.db_dir,
                                            args.out, args.target_sr)

    if failures:
        print(f"\n{len(failures)} recording(s) skipped:")
        for line in failures[:10]:
            print(f"  {line}")
        if len(failures) > 10:
            print(f"  ... and {len(failures) - 10} more")

    print(f"\nwrote {written} row(s) -> {args.out}")
    print("next: python train_grbas.py --mode finetune")


if __name__ == "__main__":
    main()
